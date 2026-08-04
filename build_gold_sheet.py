#!/usr/bin/env python3
"""
Build a BLIND human gold-coding sheet (xlsx) for the reliability subsample / Cohen's kappa.
Stratified random by journal, with narrative-form candidates oversampled (codebook priority).
Contains NO model codes — empty coding columns for the human rater to fill from the full text.
Output: data/gold_coding_sheet.xlsx  (sheet 'coding' + sheet 'Anleitung')
"""
import csv
import random
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

N_TOTAL = 80
N_NARRATIVE = 24          # oversample likely narrative-form cases
SEED = 7

NARR_CUE = re.compile(r"(?i)\bbaseline\b.{0,120}?\bno (?:statistically )?significant"
                      r"|\bno (?:statistically )?significant.{0,120}?\bbaseline\b")

have = {p.stem for p in Path("pdftext").glob("*.txt")}
corpus = [r for r in csv.DictReader(open("data/corpus.csv")) if r["record_id"] in have]

# tag narrative-likely (text contains a between-group baseline no-difference statement)
narr = []
for r in corpus:
    t = Path(f"pdftext/{r['record_id']}.txt").read_text(errors="ignore")
    if NARR_CUE.search(t):
        narr.append(r)

rng = random.Random(SEED)
rng.shuffle(narr)
pick_narr = narr[:N_NARRATIVE]
picked_ids = {r["record_id"] for r in pick_narr}

# stratified random by journal for the remainder
rest = [r for r in corpus if r["record_id"] not in picked_ids]
by_j = {}
for r in rest:
    by_j.setdefault(r["journal"], []).append(r)
need = N_TOTAL - len(pick_narr)
total_rest = len(rest)
pick_rand = []
for j, rows in by_j.items():
    rng.shuffle(rows)
    k = max(1, round(need * len(rows) / total_rest))
    pick_rand.extend(rows[:k])
rng.shuffle(pick_rand)
pick_rand = pick_rand[:need]

rows_out = ([(r, "narrative-candidate") for r in pick_narr]
            + [(r, "stratified-random") for r in pick_rand])
rng.shuffle(rows_out)

wb = Workbook()
ws = wb.active
ws.title = "coding"
ref_cols = ["record_id", "journal", "year", "doi", "title", "sample_stratum"]
code_cols = ["eligible", "exclude_reason", "subdiscipline", "n_randomised",
             "baseline_testing", "testing_form", "n_vars_tested", "n_vars_significant",
             "downstream_response", "statistician_author", "notes"]
header = ref_cols + code_cols
ws.append(header)
for c in range(1, len(header) + 1):
    cell = ws.cell(1, c)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="D9E1F2" if c <= len(ref_cols) else "FCE4D6")
for r, stratum in rows_out:
    ws.append([r["record_id"], r["journal"], r["year"], r["doi"], r["title"], stratum]
              + [""] * len(code_cols))
ws.freeze_panes = "A2"
widths = {"record_id": 16, "journal": 16, "year": 6, "doi": 26, "title": 60,
          "sample_stratum": 18}
for i, col in enumerate(header, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(col, 16)

# instructions sheet
a = wb.create_sheet("Anleitung")
guide = [
    ("BLIND CODIEREN", "Volltext in pdftext/<record_id>.txt (oder das PDF) lesen, dann Felder füllen. NICHT die Modell-Codes ansehen."),
    ("", ""),
    ("eligible", "TRUE/FALSE. FALSE -> exclude_reason setzen, Rest leer lassen."),
    ("exclude_reason", "review / protocol / secondary / design (crossover|cluster) / non-physio / no-baseline"),
    ("subdiscipline", "msk / neuro / cardioresp / other"),
    ("n_randomised", "Gesamt randomisiert (Zahl)"),
    ("baseline_testing  (PRIMÄR)", "TRUE wenn between-arm Test auf Baseline-Variablen berichtet — (a) tabellarisch p-Werte ODER (b) narrativ ('no significant differences at baseline'). Beides zählt!"),
    ("testing_form", "tabular / narrative / both / none"),
    ("n_vars_tested", "Anzahl getesteter Baseline-Variablen (p-Werte zählen; narrativ: genannte Variablen; leer wenn unzählbar)"),
    ("n_vars_significant", "Anzahl mit p<0.05"),
    ("downstream_response", "nur wenn n_vars_significant>=1: none / prespecified_adjustment / selective_posthoc_adjustment / narrative_rand_failed / other ; sonst: na_no_sig"),
    ("statistician_author", "TRUE wenn Statistiker/Methodiker unter den Autoren (laut Rolle/Affiliation)"),
    ("notes", "wörtlicher Schnipsel der Baseline-Test-Aussage"),
    ("", ""),
    ("Kanten-Fälle", "SMD ohne Test -> FALSE (empfohlene Alternative). 'No statistical comparison at baseline' -> FALSE. p nur für Primär-Outcome-Baselinewert -> zählt als 1 Variable."),
    ("sample_stratum", "narrative-candidate = gezielt übergewichtet (Narrativ-Verdacht); stratified-random = repräsentativ. Spalte beim Codieren ignorieren."),
]
for k, v in guide:
    a.append([k, v])
a.column_dimensions["A"].width = 26
a.column_dimensions["B"].width = 110
for row in a.iter_rows():
    row[1].alignment = Alignment(wrap_text=True, vertical="top")
a.cell(1, 1).font = Font(bold=True, size=13)

Path("data").mkdir(exist_ok=True)
wb.save("data/gold_coding_sheet.xlsx")
n_narr = sum(1 for _, s in rows_out if s == "narrative-candidate")
print(f"gold sheet: {len(rows_out)} records "
      f"({n_narr} narrative-candidate + {len(rows_out)-n_narr} stratified-random)")
from collections import Counter
print("by journal:", dict(Counter(r['journal'] for r, _ in rows_out)))
print("-> data/gold_coding_sheet.xlsx")
