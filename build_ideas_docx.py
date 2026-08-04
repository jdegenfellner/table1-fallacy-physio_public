#!/usr/bin/env python3
"""Build the meta-research idea catalogue as a Word .docx in the parent folder."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

OUT = Path("/Users/degn/Library/Mobile Documents/com~apple~CloudDocs/1_ZHAW/"
           "1_Publikationen/z_LOW Hanging Fruit Publications in Applied Statistics/"
           "Meta-Research_Ideen_Physio-RCT-Korpus.docx")

VERDICT_COLORS = {
    "GAP": RGBColor(0x1B, 0x7A, 0x1B),        # green
    "PARTIAL": RGBColor(0xB8, 0x6A, 0x00),    # amber
    "EXISTS": RGBColor(0xB0, 0x1B, 0x1B),     # red
    "LAUFEND": RGBColor(0x1F, 0x4E, 0x79),    # blue
}

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text, bold=False, italic=False, size=None):
    par = doc.add_paragraph()
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return par


def bullet(text, level=0):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    par.add_run(text)
    return par


def verdict_line(verdict, note=""):
    par = doc.add_paragraph()
    r = par.add_run("VERDIKT: ")
    r.bold = True
    label = {"GAP": "GAP — klar neu", "PARTIAL": "PARTIAL — Vorarbeiten vorhanden, Nische offen",
             "EXISTS": "EXISTS — weitgehend gesättigt", "LAUFEND": "LAUFENDES HAUPTPAPER"}[verdict]
    r2 = par.add_run(label)
    r2.bold = True
    r2.font.color.rgb = VERDICT_COLORS[verdict]
    if note:
        par.add_run("  — " + note)
    return par


def idea(num, title, verdict, vnote, existing, gap, angle):
    h(f"{num}. {title}", level=2)
    verdict_line(verdict, vnote)
    p("Bestehende Literatur:", bold=True)
    for e in existing:
        bullet(e)
    pp = doc.add_paragraph()
    pp.add_run("Lücke / Neuheit: ").bold = True
    pp.add_run(gap)
    pa = doc.add_paragraph()
    pa.add_run("Möglicher Titel: ").bold = True
    r = pa.add_run(angle)
    r.italic = True


# ---------------- Titlepage ----------------
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Meta-Research-Paper-Ideen auf dem Physiotherapie-RCT-Korpus")
r.bold = True
r.font.size = Pt(20)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Ein Korpus, viele Publikationen — Ideen-Katalog mit Literatur-Existenzcheck").italic = True
d = doc.add_paragraph()
d.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_run("Stand: 17.06.2026").italic = True

# ---------------- 1. Korpus ----------------
h("1. Der Korpus als wiederverwertbares Asset", level=1)
p("Aufgebaut ist ein eingefrorener Census randomisierter kontrollierter Studien (RCTs), "
  "publiziert 2015–2024 in acht führenden Physiotherapie-/Rehabilitationszeitschriften: "
  "Physical Therapy, Clinical Rehabilitation, British Journal of Sports Medicine, Journal of "
  "Orthopaedic & Sports Physical Therapy, Archives of Physical Medicine and Rehabilitation, "
  "Physiotherapy, Journal of Physiotherapy und Musculoskeletal Science and Practice.")
p("Der teure Teil — Volltext-Akquise und Extraktion der Baseline-Tabellen — ist erledigt und "
  "wiederverwertbar. Auf derselben Datenbasis lassen sich mehrere eigenständige Meta-Research-"
  "Paper publizieren. Dieses Dokument sammelt die Kandidaten und prüft je Idee, ob es bereits "
  "eine physiotherapie-/rehaspezifische Arbeit gibt.")

# ---------------- 2. Lesehilfe ----------------
h("2. Wie dieses Dokument zu lesen ist", level=1)
bullet("GAP = klar neu; keine physio-/rehaspezifische Arbeit gefunden.")
bullet("PARTIAL = es gibt Vorarbeiten (oft 2022–2025, häufig aus der Gruppe Elkins/Moseley/"
       "Pinto bzw. Innocenti/Chiarotto, meist über PEDro); eine abgegrenzte Nische bleibt aber offen.")
bullet("EXISTS = weitgehend gesättigt; nur mit sehr scharfer Abgrenzung publizierbar.")
p("Alle Literaturangaben stammen aus Web-Recherchen (Juni 2026) und sind vor Einreichung "
  "gegen PubMed zu verifizieren. Als unsicher markierte Zitate sind eigens gekennzeichnet "
  "(siehe Abschnitt 8).", italic=True)

# ---------------- 3. Warnung Hauptpaper ----------------
h("3. Wichtig für das laufende Table-1-Paper", level=1)
p("Direkter Konkurrenzbefund, der die Neuheits-Argumentation des laufenden Hauptpapers "
  "betrifft:", bold=True)
bullet("Verhagen A, Stubbs PW, Mehta P, et al. (2022). „Comparison between 2000 and 2018 on the "
       "reporting of statistical significance and clinical relevance in physiotherapy clinical "
       "trials in six major physiotherapy journals: a meta-research design.“ BMJ Open "
       "12(1):e054875. DOI 10.1136/bmjopen-2021-054875; PMID 34980625.")
p("Verhagen berichtet bereits eine physio-spezifische Prävalenz des Baseline-Signifikanztestens: "
  "Anstieg von ~28–33 % (2000) auf 61 % (2018), allerdings nur über zwei Einzeljahre, sechs "
  "Zeitschriften und ~140 Studien, und als einen von vielen Reporting-Punkten. Abgrenzung des "
  "Hauptpapers: kontinuierliches Fenster 2015–2024, acht Zeitschriften, viel größeres N und eine "
  "Table-1-spezifische, interpretations- und konsequenzfokussierte Analyse (was tun Autoren "
  "nach einem „signifikanten“ Baseline-Unterschied?). Diese Differenzierung gehört prominent ins "
  "Manuskript.")

# ---------------- 4. Übersichtstabelle ----------------
h("4. Priorisierungs-Übersicht", level=1)
rows = [
    ("Idee", "Verdikt", "Neuheit", "Hauptkonkurrenz"),
    ("B1 Multiplizität & primäre Endpunkte", "GAP", "sehr hoch", "nur generische Methodik (BMC MRM)"),
    ("C2 Statistiker-Koautorschaft", "GAP", "hoch", "nur Allgemeinmedizin (Held 2022)"),
    ("B2 Harms-/AE-Reporting", "GAP", "hoch", "nur kondition-spezifisch (Hüft-/Knie-OA)"),
    ("A3 p-curve / Excess Significance", "PARTIAL", "hoch", "Sport-Med (JSAMS 2023)"),
    ("A2 Dichotomisierung vs. Estimation", "PARTIAL", "hoch", "Hernando 2024; Verhagen 2022"),
    ("C1 Statistik-Methoden-Landschaft", "PARTIAL", "gut", "Verhagen 2022 (nur Reporting)"),
    ("A4 MCID-Interpretation", "PARTIAL", "gut", "Gianola 2019 (nur LBP, ≤2014)"),
    ("D1 Open Science (Data/Code/OA)", "PARTIAL", "gut", "Bullock PT-Preprint 2025 (prüfen!)"),
    ("D2 Sex/Gender-Reporting & -Analyse", "PARTIAL", "mittel", "Ilhan 2025 (zuerst lesen!)"),
    ("B3 Sample size / Power", "PARTIAL", "mittel", "Hernando 2024; MSK-Power-Review 2025"),
    ("B5 Missing data / ITT vs PP", "PARTIAL", "mittel (Update)", "MSK-ITT-Review 2015"),
    ("B4 CONSORT-Adhärenz", "EXISTS", "niedrig (umlenken)", "Innocenti 2022"),
    ("B6 Registrierung / Outcome-Switching", "EXISTS", "niedrig", "Silva 2023; Pinto 2013"),
    ("E1 Spin in Abstracts", "EXISTS", "niedrig", "Kumar 2023; Stubbs 2026"),
]
table = doc.add_table(rows=1, cols=4)
table.style = "Light Grid Accent 1"
for i, c in enumerate(rows[0]):
    cell = table.rows[0].cells[i]
    cell.paragraphs[0].add_run(c).bold = True
for r_ in rows[1:]:
    cells = table.add_row().cells
    for i, c in enumerate(r_):
        cells[i].text = c

# ---------------- 5. Ideen im Detail ----------------
h("5. Ideen im Detail", level=1)

h("Gruppe A — Signifikanz & Interpretation", level=2)

idea("A0", "Table 1 Fallacy — Baseline-Signifikanztesten (laufendes Hauptpaper)", "LAUFEND",
     "Prävalenz physio bereits bei Verhagen 2022 (61 % in 2018) — scharf abgrenzen (siehe Abschnitt 3).",
     ["Verhagen 2022, BMJ Open, PMID 34980625 (physio, 2 Jahre, 6 Journals).",
      "Onkologie: „… Table 1 Fallacy“, Eur J Cancer 2023/24 (PMC11512581; 765 RCTs, ~25 %).",
      "Anästhesie/Schmerz: „The Enduring Table 1 Fallacy …“, 2025 (PMC12677328).",
      "Spinal Cord 2018, nature.com/articles/s41393-018-0203-y (Methodik-Editorial)."],
     "Über die Prävalenz hinaus: Interpretation und Konsequenzen des Baseline-Testens, "
     "kontinuierlich 2015–2024, 8 Journals, großes N.",
     "The Table 1 Fallacy persists in physiotherapy: baseline significance testing in 2015–2024 "
     "RCTs across 8 rehabilitation journals — prevalence, interpretation, and consequences.")

idea("A2", "Dichotomisierung vs. Estimation (Effektgrößen + Konfidenzintervalle)", "PARTIAL",
     "Reporting-Prävalenz bereits abgedeckt; offen ist das Interpretations-/Dichotomisierungsverhalten über die Dekade.",
     ["Hernando DF, Elkins MR, Freire APCF (2024), Braz J Phys Ther 28(3):101079, PMID 38865832 — "
      "n=100 (2021): 47 % berichten ein Between-Group-CI, nur 6 % berichten UND interpretieren ein CI.",
      "Verhagen 2022, BMJ Open (CI-/Effektgrößen-/Relevanz-Reporting, 2000 vs 2018)."],
     "Nicht ob CIs/Effektgrößen berichtet werden, sondern wie Studien ihre Schlussfolgerung rahmen "
     "(rein „signifikant/nicht signifikant“ vs. Magnitude + Unsicherheit) — als Dekadentrend, großes N, "
     "vor dem Hintergrund von ASA 2016 und der „abandon significance“-Debatte 2019.",
     "From significance to estimation? A decade (2015–2024) of effect-size and confidence-interval "
     "interpretation in physiotherapy RCTs across 8 journals.")

idea("A3", "p-Wert-Verteilung: p-curve / p-hacking / Excess Significance", "PARTIAL",
     "Deskriptives p-Wert-Reporting abgedeckt; eine physio-spezifische Verteilungsanalyse fehlt.",
     ["Verhagen 2022 (deskriptiv, kein p-curve).",
      "„The bias for statistical significance in sport and exercise medicine“, J Sci Med Sport 2023, "
      "PMID 36966124 — 18 Sport-/Bewegungsmed.-Journals, Caliper-/CI-Verteilung (nicht physio, Ratio-CIs).",
      "Hintergrund: Olsson-Collentine et al., „marginal significance“-Wording."],
     "Kein p-curve / Test auf Excess Significance auf der Verteilung der berichteten primären "
     "p-Werte aus Physio-RCTs (8 Journals, 2015–2024). Kombination aus exakt-vs-Schwelle, "
     "„marginal significance“-Wording und p-curve wäre neu.",
     "Is statistical significance manufactured in physiotherapy trials? A p-curve and reporting "
     "analysis of p-values in 10 years of RCTs across 8 rehabilitation journals (2015–2024).")

idea("A4", "MCID / klinisch relevanter Unterschied — Nutzung und Interpretation", "PARTIAL",
     "MCID-Schätzmethodik gesättigt; die Interpretation innerhalb von Physio-RCTs über alle Indikationen ist weitgehend offen.",
     ["Gianola S, Castellini G, Corbetta D, Moja L (2019), Health Qual Life Outcomes 17:127, "
      "DOI 10.1186/s12955-019-1196-8 — 42 LBP-RCTs (≤2014): nur 36 % statistisch UND klinisch relevant.",
      "Verhagen 2022 — nur 22,8 % der 2018er-Studien verknüpfen Ergebnisse explizit mit klinischer Relevanz."],
     "Feldweit und indikationsagnostisch (2015–2024, 8 Journals): Wird MCID prä-spezifiziert? Wird "
     "das CI gegen die MCID interpretiert? Wie oft wird statistische mit klinischer Signifikanz verwechselt? "
     "Geeignet, auch den Overlapping-CI-Fehlschluss zu adressieren (Differenz-CI vs. MCID-Band, nicht "
     "Einzelgruppen-CIs).",
     "Statistically significant, clinically meaningless? Use and (mis)interpretation of the MCID in "
     "physiotherapy RCTs across 8 journals, 2015–2024.")

h("Gruppe B — Reporting-Qualität & Guidelines", level=2)

idea("B1", "Multiplizität & Anzahl/Klarheit primärer Endpunkte", "GAP",
     "Stärkste greenfield-Lücke; bisher nur generische, fächerübergreifende Methodik.",
     ["Bell ML et al. (2022), BMC Med Res Methodol 22:46, PMC8818238 — 138 RCTs (allgemein), ~48 % korrigieren.",
      "„… handling of multiple outcomes …“, scoping review, PMC12289522 (Methodenliteratur).",
      "Arch Phys Med Rehabil-Kommentar „Single or Multiple Outcomes?“ (konzeptionell, keine Prävalenz)."],
     "Keine empirische Meta-Research über Physio-RCTs (2015–2024, 8 Journals) zu: Anzahl primärer "
     "Endpunkte, Definitionsklarheit, Anzahl Outcomes×Zeitpunkte, und ob Multiplizität korrigiert wurde. "
     "Direkt aus dem Korpus extrahierbar.",
     "Counting what counts: number and definitional clarity of primary outcomes and (un)corrected "
     "multiplicity in physiotherapy RCTs, 2015–2024.")

idea("B2", "Adverse-Event-/Harms-Reporting", "GAP",
     "Vorhandene Arbeiten sind kondition-spezifisch oder inzidenz-fokussiert; ein breiter Reporting-Qualitäts-Audit fehlt.",
     ["James KA, von Heideken J, Iversen MD (2021), Phys Ther 101(11):pzab195, PMID 34730830 — "
      "Hüft-OA-Übungsstudien: nur 42,9 % dokumentieren AEs.",
      "von Heideken et al. (2021), Phys Ther, Knie-OA (PMC8557852).",
      "Niemeijer A et al. (2020), Br J Sports Med 54(18):1073, PMID 31563884 — Inzidenz, nicht Reporting-Qualität.",
      "Hintergrund: CONSORT Harms 2022 (Junqueira DR et al., PMID 37100738)."],
     "Kein journal-definierter (statt kondition-definierter) Physio-Korpus zur Harms-Reporting-Qualität / "
     "CONSORT-Harms-Adhärenz. 2015–2024 erlaubt eine Vorher-Nachher-Analyse um das CONSORT-Harms-Update 2022.",
     "Are harms reported at all? AE reporting quality and CONSORT-Harms adherence in physio/exercise "
     "RCTs across eight journals (2015–2024).")

idea("B3", "Sample size / a-priori Power — Berichten und Angemessenheit", "PARTIAL",
     "Starke aktuelle Vorarbeit; eine fokussierte Power-Adäquanz-Arbeit hat noch Raum.",
     ["Hernando DF, Elkins MR, Freire APCF (2024), Braz J Phys Ther 28(3):101079 — n=100 (2021), "
      "~⅓ ohne Fallzahlplanung; 59 % berechnet UND erreicht.",
      "„Statistical Power in Musculoskeletal Research: A Meta-Review of 266 RCTs“, Sports Med Open "
      "(2025), DOI 10.1186/s40798-025-00908-8 (nur MSK; abstract-/indexbasiert prüfen)."],
     "Dekadenlang (2015–2024), 8 Journals: Qualität/Vollständigkeit der a-priori-Berechnung "
     "(Effektgrößen-Begründung, Varianzquelle, Attrition), Realismus der angenommenen vs. beobachteten "
     "Effektgröße und neu berechnete erreichte Power.",
     "Are physiotherapy RCTs powered for the effects they assume? A decade (2015–2024) of sample-size "
     "justification quality and recomputed achieved power across eight journals.")

idea("B4", "CONSORT-Adhärenz / Reporting-Vollständigkeit", "EXISTS",
     "Sehr naher Präzedenzfall; nur mit Umlenkung auf die Baseline-Nische sinnvoll.",
     ["Innocenti T, Giagio S, Salvioli S, et al. (2022), Arch Phys Med Rehabil, "
      "DOI 10.1016/j.apmr.2022.01.156, PMID 35192799 — 200 RCTs, 2011–2020, 68 Reha-Journals; "
      "mittlere CONSORT-Adhärenz 65 %.",
      "Moseley AM, Elkins MR, et al. (2014), Physiother Can, PMID 24719507 — 14.910 Reports auf der PEDro-Skala."],
     "Eine allgemeine CONSORT-Adhärenz-Arbeit würde Innocenti 2022 duplizieren. Neu nur als "
     "Umlenkung auf CONSORT-Item 15 (Baseline-Daten) und das abgeratene Baseline-Signifikanztesten — "
     "verbindet sich mit dem Table-1-Thema.",
     "Baseline data reporting and the persistence of baseline significance testing (CONSORT item 15) "
     "in RCTs from eight leading physiotherapy journals, 2015–2024.")

idea("B5", "Missing Data & Analysepopulation (ITT vs. Per-Protocol)", "PARTIAL",
     "Naher MSK-Review vorhanden, aber ~10 Jahre alt — als „decade-on“-Update positionieren.",
     ["„… variable use of the intention-to-treat principle in musculoskeletal RCTs with missing data“, "
      "J Clin Epidemiol 2015, PMID 25304501 — 91 MSK-RCTs, 2010–2011; 38 % volle ITT; LOCF dominant.",
      "Meta-epidemiologisch: Attrition/Missing/ITT vs. Effektgrößen-Bias (393 Studien)."],
     "2015–2024, alle 8 Journals (breiter als MSK — inkl. Neuro/Kardioresp via Clin Rehabil, Arch Phys "
     "Med Rehabil): Zeittrend und Verschiebung von LOCF zu Mixed Models / Multipler Imputation.",
     "A decade on: ITT vs per-protocol and the move from LOCF to mixed models — missing-data handling "
     "in physiotherapy RCTs, 2015–2024.")

idea("B6", "Trial-Registrierung & Outcome-Switching", "EXISTS",
     "Am stärksten gesättigt (kanonische Physio-Autoren); nur sehr enge Nische offen.",
     ["Silva NS, Elkins MR, Lemes ÍR, Stubbs PW, Franco MR, Pinto RZ (2023), Musculoskelet Sci Pract "
      "67:102854, PMID 37657398 — 200 Studien (2019 vs 2009); prospektive Registrierung 18 % vs 6 %; "
      "selektives Outcome-Reporting 73 %.",
      "Pinto RZ, Elkins MR, Moseley AM, et al. (2013), Phys Ther 93(3):299 — Basisdatensatz 2009."],
     "Eng. Neu nur als jährlicher Trend 2015–2024 oder als Richtung/Effekt der Outcome-Switches "
     "(begünstigen sie signifikante Ergebnisse?) statt reiner Diskrepanzzählung.",
     "Did journal registration policies work? Annual trends in prospective registration and the "
     "direction of outcome switching in physiotherapy RCTs across eight journals, 2015–2024.")

h("Gruppe C — Methoden & Menschen", level=2)

idea("C1", "Statistik-Methoden-Landschaft & Trends", "PARTIAL",
     "Physio-Statistik-Meta-Research deckt nur Reporting-Scheiben ab, nicht Modelle/Software.",
     ["Verhagen 2022, BMJ Open (nur Reporting, 2 Jahre, endet 2018).",
      "Roush JR et al. (2015), J Phys Ther Educ 29(1):5 — 2009–2010, keine Software/Mixed/Bayes.",
      "Tilson JK et al. (2016), BMC Med Educ 16:118, PMID 27101814 — Begriffsfrequenz, 2011–2012.",
      "Policy: ISPJE „estimation over NHST“, Phys Ther 2022;102(6):pzac066."],
     "Keine Arbeit charakterisiert die Modell-/Test-Landschaft (t-Test/ANOVA/ANCOVA/Mixed/GEE), die "
     "Verbreitung von Mixed Models und Bayes sowie die Software (SPSS/R/Stata/Stan) als kontinuierlichen "
     "Trend 2015–2024.",
     "From NHST to mixed models and Bayes? A decade (2015–2024) of statistical methods, modeling "
     "choices, and software in physio/rehab RCTs across eight journals.")

idea("C2", "Statistiker-Koautorschaft und Reporting-Qualität", "GAP",
     "In der Allgemeinmedizin gut belegt, in Physio/Reha nie gemacht — sauberste GAP.",
     ["Held U, Steigmiller K, Hediger M, et al. (2022), PLoS ONE 17(3):e0264819, PMC8896706 — "
      "Biostatistiker-Koautorschaft → +1,60 Reporting-Qualitätseinheiten.",
      "Diaz-Ordaz K, Froud R, Sheehan B, Eldridge S (2013), BMC Med Res Methodol 13:127, PMID 24148859.",
      "Sainani KL, Borg DN, Caldwell AR, et al. (2020), Br J Sports Med, PMID 32816788 — nur ~13,3 % "
      "der Arbeiten mit Statistik-Koautor (Basisrate, kein Outcome-Link)."],
     "Der Exposure→Outcome-Test (Statistiker-Koautor → Reporting-/Statistikqualität) wurde in "
     "Physio/Reha nie durchgeführt. Schlüsselentscheidung: Operationalisierung der „Statistiker-"
     "Beteiligung“ (Affiliation vs. Abschluss vs. benannte Rolle).",
     "Statistician co-authorship and reporting quality of physio/rehab RCTs: a meta-research study "
     "of 8 journals (2015–2024).")

h("Gruppe D — Transparenz & Gerechtigkeit", level=2)

idea("D1", "Open-Science-Praktiken (Data-/Code-Sharing, Open Access)", "PARTIAL",
     "Aktiv beforscht in Nachbarfeldern; kein exakter Treffer — ein PT-Preprint ist aber zu prüfen.",
     ["Bullock GS, Ward P, Impellizzeri FM, et al. (2023), J Orthop Sports Phys Ther 53(12):735, "
      "PMID 37860866 — Sportmed.-Journals, alle Studientypen, nur 2022.",
      "Elghzali A et al. (2025), BMC Med Res Methodol 25:129, PMID 40348959 — 5 Reha-Journals, "
      "nur Data-Sharing-Statements, 2017–2023.",
      "KONKURRENZRISIKO: Bullock GS et al. (2025, medRxiv 2025.09.23.25336497), "
      "„Replicability and transparency in physical therapy research“ — Scope unklar, ZUERST LESEN."],
     "Keine Arbeit kombiniert RCT-only + diese 8 Journals + alle drei Dimensionen (Data + Code + OA) "
     "+ vollen Trend 2015–2024.",
     "From statements to substance: a decade (2015–2024) of data-sharing, code-sharing, and open-access "
     "in RCTs across eight physiotherapy/rehabilitation journals.")

idea("D2", "Sex/Gender-Reporting und sexbasierte Analyse", "PARTIAL",
     "Sehr naher Konkurrent; Nische enger, aber real (Analyse-Validität).",
     ["KONKURRENT (zuerst lesen): Ilhan E, Solis K, Liu C, et al. (2025), Physiotherapy, "
      "DOI 10.1016/j.physio.2024.101450, PMID 39689410 — 182 RCTs, 2018–2024; nur Reporting/Terminologie.",
      "Wiley E, Noguchi KS, … Tang A (2024), PLoS One, PMID 39383122 — Enrollment + Interaktion, nur Stroke.",
      "Welch V, Doull M, et al. (2017), Res Integr Peer Rev, PMID 29451565 — kanonische Basis.",
      "SAGER-Guidelines: Heidari S, et al. (2016), DOI 10.1186/s41073-016-0007-6."],
     "Kombiniert weder jemand Frauen-Enrollment-Adäquanz (vs. Krankheitsprävalenz) MIT Sex als "
     "Analysevariable über den breiten 8-Journal-Korpus ab 2015. Differenzierer (passt zu den "
     "Statistik-Regeln): Wurden Sex-Unterschiede über valide Interaktionstests belegt — oder über "
     "ungültiges subgruppenweises Testen / Overlapping-CI-Schlüsse?",
     "Counted but not analyzed: female enrollment and the (mis)use of sex as an analytic variable in "
     "physio/rehab RCTs, 2015–2024.")

h("Gruppe E — Gesättigt / niedrige Priorität", level=2)

idea("E1", "Spin in Abstracts (nicht-signifikante primäre Endpunkte)", "EXISTS",
     "Mehrere physio-spezifische Spin-Arbeiten existieren; „what’s new?“-Risiko hoch.",
     ["Kumar JS, Prakash V, Hariohm K, Kumar IN (2023), Physiother Pract Res 44(2):109, "
      "DOI 10.3233/PPR-230747 — 44 Studien; Spin in 90 % der Abstracts.",
      "Stubbs PW et al. (2026), Cochrane Evidence Synthesis and Methods, DOI 10.1002/cesm.70072 — "
      "160 Studien (Robotik); 97 % mit Spin.",
      "Patterson et al. (2021), PMID 34261764 — RCT zur Leserwahrnehmung (anderes Design)."],
     "Nur die breite, journalübergreifende repräsentative Stichprobe (8 Top-Journals, 2015–2024) ist "
     "offen; liest sich aber inkrementell.",
     "Spin in the abstracts of physiotherapy RCTs with non-significant primary outcomes: a 2015–2024 "
     "cross-journal meta-research study of 8 leading rehabilitation journals.")

# ---------------- 6. Empfehlung ----------------
h("6. Empfehlung", level=1)
bullet("Sauberste neue Paper: B1 (Multiplizität & primäre Endpunkte) und C2 (Statistiker-Koautorschaft) "
       "— echte GAPs, direkt aus dem Korpus bedienbar.")
bullet("Methodisch markantestes Paper: A3 (p-curve / Excess Significance).")
bullet("Gut publizierbar mit Abgrenzung: A2 (Estimation-Denken), C1 (Methoden-Landschaft), "
       "A4 (MCID-Interpretation), B2 (Harms), D1 (Open Science).")
bullet("Nur als explizites Update/Replikation: B3 (Power), B5 (ITT/Missing), D2 (Sex/Gender) — jeweils "
       "den genannten Konkurrenten zuerst lesen.")
bullet("Niedrige Priorität / gesättigt: B4 (CONSORT, außer Baseline-Umlenkung), B6 (Registrierung), E1 (Spin).")
p("Roter Faden des Programms: Der komparative Vorteil dieses Korpus sind die extrahierten "
  "Baseline-Tabellen und der volle Dekaden-Verlauf — das nutzen die bestehenden Konkurrenten nicht. "
  "Ideen, die genau das ausspielen, haben die beste Publikationschance.", italic=True)

# ---------------- 7. Methodik des Checks ----------------
h("7. Methodik des Literatur-Checks", level=1)
p("Vier parallele Recherche-Durchläufe (Juni 2026) über vier Themencluster, jeweils mit Websuche "
  "und Verifikation über PubMed/PMC/Verlagsseiten, wo möglich. Verdikte spiegeln den Stand dieser "
  "Recherche; sie ersetzen keine vollständige systematische Suche.")

# ---------------- 8. Caveats ----------------
h("8. Vorbehalte / zu verifizierende Zitate", level=1)
p("Vor Einreichung direkt gegen PubMed prüfen — folgende Angaben waren nur indirekt belegbar:")
bullet("Bullock GS et al. (2025) PT-Preprint (medRxiv 2025.09.23.25336497): genauer Scope (RCT-only? "
       "Journals? Code/OA-Trends?) ungeprüft — größtes Überlappungsrisiko für D1.")
bullet("„Statistical Power in Musculoskeletal Research“ (Sports Med Open 2025): n und Befunde nur "
       "abstract-/indexbasiert.")
bullet("Genaue Autorenreihenfolge der Onkologie-Table-1-Arbeit (Eur J Cancer) und die generische "
       "„Table I Fallacy: P values …“-Preprint-Angabe.")
bullet("„Ann Rehabil Med 2017“-Trends-Arbeit (Byline ungeprüft) sowie ein etwaiges ADVANCE-"
       "Ergebnispaper zur AE-Berichterstattung (B2).")
bullet("Eine in einer KI-Zusammenfassung aufgetauchte „Borg et al. 2023, excess significance in "
       "sports therapy and rehabilitation“ konnte NICHT als eigenständige Arbeit verifiziert werden — "
       "nicht ohne Bestätigung zitieren.")

doc.save(OUT)
print("saved:", OUT)
print("paragraphs:", len(doc.paragraphs))
