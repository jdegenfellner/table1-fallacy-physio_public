#!/usr/bin/env python3
"""Build the 'accessibility / open access' paper-idea concept note as a Word .docx
in the parent Low Hanging Fruit program folder."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

OUT = Path("/Users/degn/Library/Mobile Documents/com~apple~CloudDocs/1_ZHAW/"
           "1_Publikationen/z_LOW Hanging Fruit Publications in Applied Statistics/"
           "Paper-Idee_Zugaenglichkeit_Physio-RCT-Korpus.docx")

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def h(t, lvl=1):
    doc.add_heading(t, level=lvl)


def p(t, bold=False, italic=False):
    par = doc.add_paragraph()
    r = par.add_run(t); r.bold = bold; r.italic = italic
    return par


def bullet(t, lvl=0):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.left_indent = Inches(0.25 + 0.25 * lvl)
    par.add_run(t)


# title
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Wer kommt an die Evidenz? Zugänglichkeit und Paywalls im physiotherapeutischen RCT-Korpus")
r.bold = True; r.font.size = Pt(18)
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
s.add_run("Paper-Idee (eigenständig) auf demselben Korpus — Low-Hanging-Fruit-Programm").italic = True
d = doc.add_paragraph(); d.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_run("Stand: 23.06.2026").italic = True

h("1. Motivation", 1)
p("Öffentlich finanzierte klinische Forschung wird von Steuerzahlern bezahlt, von Autoren "
  "unentgeltlich geschrieben und begutachtet — und anschließend von kommerziellen Verlagen "
  "hinter Paywalls verkauft, oft mit Gewinnmargen von 30–40 %. Für die evidenzbasierte "
  "Physiotherapie heißt das konkret: Wer behandelt, kommt an einen relevanten Teil der "
  "eigenen Evidenzbasis nur über kostenpflichtige Einzelkanäle. Wie groß dieser Teil ist, wie "
  "er sich über die Zeit und über Verlage verteilt, ist selten quantifiziert worden — und lässt "
  "sich aus unserem bereits aufgebauten Beschaffungs-Funnel direkt messen.")

h("2. Forschungsfragen", 1)
bullet("Welcher Anteil der physiotherapeutischen RCTs (8 führende Journale, 2015–2024) ist "
       "frei zugänglich (Gold/Green/Hybrid/Bronze-OA) vs. ausschließlich hinter Paywall?")
bullet("Über welchen Kanal wird ein Volltext tatsächlich beschaffbar (PMC-OA, Repositorium/Green, "
       "Verlags-TDM, institutionelle Lizenz, Fernleihe, gar nicht) — und mit welchem Aufwand/Kosten?")
bullet("Wie verteilt sich (Un-)Zugänglichkeit über Verlage und Journale? Gibt es eine "
       "Konzentration bei einzelnen Verlagen?")
bullet("Trend über die Dekade: Wird neuere Forschung zugänglicher (OA-Mandate, Plan S) oder bleibt "
       "der jüngste Jahrgang am stärksten verschlossen?")
bullet("Kontrast: deckt sich „offiziell als OA deklariert“ (Unpaywall) mit „praktisch in einem "
       "Bulk-/TDM-Workflow beschaffbar“?")

h("3. Die Daten liegen bereits vor", 1)
p("Der teure Teil ist erledigt: Für den eingefrorenen Census wurde der Volltext über eine "
  "Kaskade legaler Kanäle beschafft und dabei protokolliert, welcher Kanal griff. Dieser "
  "Beschaffungs-Funnel IST der Datensatz für dieses Paper — kein neuer Erhebungsaufwand.")

h("4. Vorläufige Kennzahlen (aus dem laufenden Projekt, provisorisch)", 1)
bullet("Census: 1663 Records; Analyse-Frame needs_pdf: 1556 (8 Journale, 2015–2024).")
bullet("Frei via PMC (Green-OA-Proxy, PMCID vorhanden): 269 = 16 %. Die übrigen ~84 % erforderten "
       "Verlags-TDM, institutionelle Lizenz oder manuelle Beschaffung.")
bullet("Volltext-Abdeckung aktuell: 1474/1556 = 95 % — erreicht erst durch gestapelte Kanäle "
       "(PMC-OA, Elsevier-TDM, Green-OA/Unpaywall, manuell; SAGE-TDM ausstehend).")
bullet("Abdeckung pro Journal: 7 von 8 bei ~100 % (Arch Phys Med Rehabil, BJSM, JOSPT, J Physiother, "
       "MSK Sci Pract, Phys Ther, Physiotherapy) — Clinical Rehabilitation (SAGE) nur 85 %.")
bullet("Der gesamte verbleibende Rest (82 Artikel) stammt aus EINEM Verlag (SAGE/Clinical "
       "Rehabilitation), und 79 der 82 sind 2022 oder neuer. Single-Publisher-Konzentration + "
       "Aktualitäts-Paywall in Reinform.")
p("Hinweis: Diese Zahlen sind ein Nebenprodukt der laufenden Beschaffung und vor einer Publikation "
  "sauber zu re-klassifizieren (siehe Methoden).", italic=True)

h("5. Methoden (Skizze)", 1)
bullet("Pro Record den OA-Status über Unpaywall abfragen und kategorisieren: closed / bronze / "
       "hybrid / green / gold (Diamond separat, falls erkennbar).")
bullet("Pro Record den real erfolgreichen Beschaffungskanal codieren (PMC-OA / Repositorium / "
       "Verlags-TDM / institutionelle Lizenz / Fernleihe / nicht beschaffbar) inkl. Aufwand.")
bullet("Deskriptiv: Anteile gesamt, nach Verlag/Journal, nach Jahr (Zeittrend). Prävalenzen mit "
       "Wilson-CI; Trend per logistischer Regression (OA ja/nein ~ Jahr).")
bullet("Inferenz-Guardrail: Unterschiede zwischen Verlagen/Jahren über die Differenz selbst "
       "(Regressionskoeffizienten + CI) bewerten — NICHT über die Überlappung der Gruppen-CIs.")
bullet("Kosten-/Aufwandsschätzung für die irreduzible Paywall-Residue (Fernleihe/Einzelkauf je "
       "Artikel) — was würde es kosten, EINEN Querschnitt der eigenen Evidenzbasis vollständig "
       "zu lesen?")

h("6. Neuheit / Positionierung", 1)
bullet("OA-Prävalenz-Studien gibt es feldübergreifend (z. B. Piwowar et al. 2018, „State of OA“). "
       "Physiotherapie-/Reha-spezifisch und auf einen RCT-Census fokussiert ist es kaum gemacht — "
       "und der Kontrast „deklariertes OA vs. tatsächliche Beschaffbarkeit im TDM-Workflow“ ist neu.")
bullet("Reflexiver Reiz: ein Meta-Research-Paper über die Zugänglichkeit genau jener Literatur, "
       "die ein anderes Paper desselben Korpus methodisch auswertet — die Zugangsbarriere als "
       "Forschungsgegenstand statt nur als Hindernis.")
bullet("Policy-Relevanz: belegt die Paywall-Konzentration quantitativ; anschlussfähig an Plan S, "
       "OA-Mandate, Big-Deal-Verhandlungen der Bibliothekskonsortien.")

h("7. Direkt verwertbar: Limitations-Absatz für das Table-1-Paper", 1)
p("Fertiger Baustein (sachlich formuliert) für die Limitations/Methods des laufenden Papers:")
q = doc.add_paragraph()
q.paragraph_format.left_indent = Inches(0.3)
qr = q.add_run(
    "Full-text retrieval relied on a cascade of legal channels (PMC open access, publisher "
    "text-and-data-mining APIs, repository/green open-access copies, institutional access, and "
    "manual retrieval). Only 16% of records had a freely available PMC copy; the remainder "
    "required licensed or mining-based access. We obtained full text for 95% of the eligible "
    "frame (1474/1556). The residual 5% were concentrated entirely in a single subscription "
    "journal (Clinical Rehabilitation) and overwhelmingly in the most recent years (79/82 from "
    "2022 onward), reflecting a single-publisher, recency-weighted paywall barrier rather than "
    "random missingness. Because this residue is non-random (one journal), we report a "
    "sensitivity analysis bounding the prevalence over best/worst-case coding of the "
    "unretrieved records.")
qr.italic = True

h("8. Vorbehalte", 1)
bullet("Vorläufige Zahlen vor Publikation re-klassifizieren (Unpaywall-Status, sauberer "
       "Kanal-Code statt PMCID-Proxy).")
bullet("PMCID ist nur ein grober Green-OA-Proxy; einige OA-Artikel haben keine PMCID, manche "
       "PMCID-Records sind embargoiert.")
bullet("Der protokollierte Kanal sagt, was zuERST griff, nicht zwingend den günstigsten möglichen.")

doc.save(OUT)
print("saved:", OUT)
