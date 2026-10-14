#!/usr/bin/env python3
"""
Build the BLIND human validation sheet for the analytic-response classification
(all trials with >= 1 "significant" baseline variable). Revision 1, Oct 2026.

Reads BMC_MRM_revision1/response_trials_key_DO_NOT_OPEN_WHILE_CODING.csv (written by
revision_analysis.py), drops the model's category, keeps the verbatim evidence snippet
and adds an automatic text excerpt (sentences near baseline/significant/randomisation
cues) from pdftext/<record_id>.txt. Output:
  BMC_MRM_revision1/validation_response/RESPONSE_VALIDATION_SHEET.xlsx
"""
import csv, re
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

SRC = Path('BMC_MRM_revision1/response_trials_key_DO_NOT_OPEN_WHILE_CODING.csv')
OUT = Path('BMC_MRM_revision1/validation_response'); OUT.mkdir(exist_ok=True)
corpus = {r['record_id']: r for r in csv.DictReader(open('data/corpus.csv'))}

CUE = re.compile(r"(?i)(baseline|randomi[sz]|imbalanc|unequal|not (?:well )?balanced|despite|covariate|adjust|confound|"
                 r"significant(?:ly)? (?:differ|higher|lower|more|less|older|younger)|p\s*[<=]\s*0?\.0[0-4])")
SENT = re.compile(r'(?<=[.!?])\s+(?=[A-Z(])')

def excerpt(rid, limit=2500):
    p = Path(f'pdftext/{rid}.txt')
    if not p.exists(): return '(no local text; use the PDF)'
    t = p.read_text(errors='ignore')
    t = re.sub(r'\s+', ' ', t)
    t = ILLEGAL_CHARACTERS_RE.sub('', t)
    sents = SENT.split(t)
    keep = []
    for i, s in enumerate(sents):
        if CUE.search(s) and 30 < len(s) < 600:
            keep.append(s.strip())
    out = ' || '.join(keep)
    return out[:limit] + (' [...]' if len(out) > limit else '')

rows = list(csv.DictReader(open(SRC)))
wb = Workbook(); ws = wb.active; ws.title = 'coding'
ref = ['Nr', 'record_id', 'journal', 'year', 'doi', 'title', 'pdf', 'n_vars_tested', 'n_vars_significant_model_count',
       'evidence_snippet', 'auto_excerpt']
code = ['n_vars_significant_human', 'response_human', 'adjustment_basis', 'wording_quote', 'notes']
hdr = ref + code
ws.append(hdr)
for c in range(1, len(hdr)+1):
    cell = ws.cell(1, c); cell.font = Font(bold=True)
    cell.fill = PatternFill('solid', fgColor='D9E1F2' if c <= len(ref) else 'FCE4D6')
for i, r in enumerate(rows, 1):
    rid = r['record_id']
    ws.append([i, rid, r['journal'], r['year'], r['doi'], corpus.get(rid, {}).get('title', ''),
               f"pdfs/{rid}.pdf" if r['has_pdf'] == 'True' else '(no PDF archived; open via DOI)',
               r['n_vars_tested'], r['n_vars_significant'], ILLEGAL_CHARACTERS_RE.sub('', r['evidence_snippet']), excerpt(rid)] + ['']*len(code))
ws.freeze_panes = 'C2'
for col, wdt in {'A': 5, 'B': 16, 'C': 14, 'D': 6, 'E': 24, 'F': 50, 'G': 22, 'H': 8, 'I': 10, 'J': 60, 'K': 80,
                 'L': 10, 'M': 24, 'N': 16, 'O': 40, 'P': 30}.items():
    ws.column_dimensions[col].width = wdt
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')

a = wb.create_sheet('Anleitung')
guide = [
    ('BLIND CODIEREN', 'Pro Trial den Volltext (PDF oder pdftext) lesen. Die Modell-Kategorie ist NICHT im Sheet. '
                       'evidence_snippet und auto_excerpt sind nur Lesehilfen (wörtliche Textstellen), kein Urteil.'),
    ('', ''),
    ('n_vars_significant_human', 'Anzahl Baseline-Variablen, die der Artikel als "significant" (p<0.05) berichtet. 0 wenn keine. '
                                 'Falls 0: response_human = na_no_sig.'),
    ('response_human', 'GENAU EINE Kategorie (Vorrang bei Mehrfachtreffern: selective_posthoc_adjustment > results_discounted > randomisation_doubted > neutral_mention > none):\n'
                       '  none = Differenz wird nirgends im Text kommentiert\n'
                       '  neutral_mention = Differenz wird erwähnt, ohne Zweifel an der Randomisierung und ohne Konsequenz für die Analyse\n'
                       '  randomisation_doubted = Text schreibt die Differenz einer fehlgeschlagenen/unzureichenden Randomisierung zu oder nennt sie eine Limitation der Randomisierung\n'
                       '  results_discounted = Ergebnisse werden wegen der Differenz relativiert / als möglicherweise verzerrt bezeichnet (ohne Adjustierung)\n'
                       '  selective_posthoc_adjustment = Adjustierung/Kovariate wurde WEGEN des Baseline-Tests gewählt (nach Sicht der Daten)\n'
                       '  prespecified_adjustment = Adjustierung war laut Text/Protokoll vorab festgelegt\n'
                       '  other = passt in keine Kategorie (notes!)'),
    ('adjustment_basis', 'Nur wenn adjustiert wurde: significance (wegen p-Wert) / clinical (wegen klinisch relevanter Grösse) / prespecified / unclear'),
    ('wording_quote', 'Wörtliches Zitat des Satzes, der die Einordnung trägt (z. B. "randomisation failed to ...").'),
    ('notes', 'Freitext.'),
    ('', ''),
    ('Mapping auf die alten Kategorien', 'narrative_rand_failed (alt) = randomisation_doubted ODER results_discounted; '
                                        'neutral_mention war bisher nicht getrennt (zählt neu NICHT als inappropriate).'),
]
for k, v in guide: a.append([k, v])
a.column_dimensions['A'].width = 30; a.column_dimensions['B'].width = 120
for row in a.iter_rows():
    row[1].alignment = Alignment(wrap_text=True, vertical='top')
a.cell(1, 1).font = Font(bold=True, size=13)
wb.save(OUT/'RESPONSE_VALIDATION_SHEET.xlsx')
print(f"wrote {OUT/'RESPONSE_VALIDATION_SHEET.xlsx'} with {len(rows)} trials; "
      f"{sum(1 for r in rows if r['has_pdf']=='True')} with archived PDF, {sum(1 for r in rows if r['has_text']=='True')} with local text")
